"""plans/116 §10.3 후속 — 양식 채우기 결함 2건 회귀.

1. F-SERVER-LIST: 서버목록 양식의 IP주소가 50대 중 8대 비었다. 폼필 피벗 SQL(코드 조립)이
   `IP주소 → EAV:IPaddress` 매핑을 그대로 EAV 속성으로 뽑았는데, 프로필 `value_joins` 가
   「EAV IPaddress = 엔티티 ipaddress 컬럼」이라 선언하고 실제 값은 직접 컬럼에만 있었다.
   `value_joins` 로 선언된 EAV 속성은 엔티티 직접 컬럼으로 뽑는다.
2. F-WORD: 보고서 docx 의 `{{부서}}`·`{{서버수}}` 가 빈칸이었다. 매핑 불가 자리 표시는
   조회 행에 값이 없으면 비웠다. 질의에 명시한 값(「부서는 X」)과 행 수를 결정적으로 채운다.
"""

from __future__ import annotations

import io

from docx import Document

from src.db_adapters.polestar.assembler import build_form_fill_pivot_sql
from src.document.word_writer import fill_word_template, resolve_placeholder_literals

_PATTERN = {
    "entity_table": "cmm_resource",
    "config_table": "core_config_prop",
    "attribute_column": "name",
    "value_column": "stringvalue_short",
    "value_joins": [
        {"eav_attribute": "Hostname", "eav_value_column": "stringvalue_short",
         "entity_column": "hostname"},
        {"eav_attribute": "IPaddress", "eav_value_column": "stringvalue_short",
         "entity_column": "ipaddress"},
    ],
    "direct_join": {"entity_column": "resource_conf_id", "config_column": "configuration_id"},
}


# ── 1. IP주소 출처 ────────────────────────────────────────────────


def test_value_join_attributes_use_entity_columns() -> None:
    sql = build_form_fill_pivot_sql(
        [("상태", "cmm_resource.avail_status")],
        [("제조사", "Vendor"), ("호스트명", "Hostname"), ("IP주소", "IPaddress")],
        [("CPU코어수", "LOGICALCORE", "server.Cpus")],
        _PATTERN,
    )
    assert "THEN c.ipaddress END) AS \"IP주소\"" in sql
    assert "THEN c.hostname END) AS \"호스트명\"" in sql
    assert "cc.name='IPaddress'" not in sql
    assert "cc.name='Hostname'" not in sql
    # 선언이 없는 속성은 종전대로 EAV
    assert "cc.name='Vendor'" in sql
    # 양식 열 순서 보존(SELECT 순서 = 매핑 순서)
    assert sql.index('"제조사"') < sql.index('"호스트명"') < sql.index('"IP주소"')


def test_without_value_joins_eav_is_kept() -> None:
    pattern = {k: v for k, v in _PATTERN.items() if k != "value_joins"}
    sql = build_form_fill_pivot_sql([], [("IP주소", "IPaddress")], [], pattern)
    assert "cc.name='IPaddress'" in sql


# ── 2. Word 자리 표시 ─────────────────────────────────────────────


def _docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "호스트명"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_QUERY = "첨부한 보고서 양식에 전체 서버 현황을 넣어서 파일로 만들어줘. 부서는 인프라운영팀"
_MAPPING = {"호스트명": "EAV:Hostname", "부서": None, "서버수": None}
_ROWS = [{"호스트명": f"svr-{i}"} for i in range(54)]


def test_resolve_placeholder_literals_from_query_and_row_count() -> None:
    lits = resolve_placeholder_literals(["부서", "서버수"], _MAPPING, _ROWS, _QUERY)
    assert lits == {"부서": "인프라운영팀", "서버수": "54"}


def test_resolve_skips_mapped_and_unknown_placeholders() -> None:
    lits = resolve_placeholder_literals(
        ["호스트명", "담당자"], {"호스트명": "EAV:Hostname", "담당자": None}, _ROWS, _QUERY,
    )
    assert lits == {}


def test_query_value_stops_at_delimiters() -> None:
    lits = resolve_placeholder_literals(
        ["부서"], {"부서": None}, [], "부서: 인프라운영팀, 기간은 지난달로 해줘",
    )
    assert lits == {"부서": "인프라운영팀"}


def test_fill_word_template_writes_literals_into_body() -> None:
    file_data = _docx(["작성 부서: {{부서}}", "대상 서버 수: {{서버수}}"])
    template = {"placeholders": ["부서", "서버수"], "tables": []}
    out = fill_word_template(
        file_data, template, _MAPPING, _ROWS,
        literal_values={"부서": "인프라운영팀", "서버수": "54"},
    )
    texts = [p.text for p in Document(io.BytesIO(out)).paragraphs]
    assert "작성 부서: 인프라운영팀" in texts
    assert "대상 서버 수: 54" in texts


def test_row_value_wins_over_literal() -> None:
    file_data = _docx(["작성 부서: {{부서}}"])
    out = fill_word_template(
        file_data, {"tables": []}, {"부서": None}, [{"부서": "행 값"}],
        literal_values={"부서": "질의 값"},
    )
    assert "작성 부서: 행 값" in [p.text for p in Document(io.BytesIO(out)).paragraphs]
