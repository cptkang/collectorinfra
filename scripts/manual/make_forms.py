"""사례용 양식 픽스처 생성 (plans/116 §4.9.2 6장). 기존 테스트 템플릿을 복사하고 Word 양식을 만든다."""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
from docx import Document

REPO = Path(__file__).resolve().parents[2]
OUT = REPO / "scripts" / "manual" / "fixtures" / "forms"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPO / "testdata/templates/server_list_template.xlsx", OUT / "서버목록_양식.xlsx")
    shutil.copy2(
        REPO / "testdata/scenarios/fixtures/adhoc_server_info.xlsx", OUT / "서버정보_양식.xlsx"
    )
    shutil.copy2(
        REPO / "testdata/scenarios/fixtures/adhoc_cpu_memory.xlsx", OUT / "성능요약_양식.xlsx"
    )
    # 역질문(U-29) 사례용 — DB 에 없는 열(담당자·연락처)을 둔다
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "서버담당"
    ws.append(["호스트명", "OS종류", "담당자", "담당자 연락처"])
    wb.save(OUT / "서버담당_양식.xlsx")
    doc = Document()
    doc.add_heading("서버 현황 보고서", level=1)
    doc.add_paragraph("작성 부서: {{부서}}")
    doc.add_paragraph("대상 서버 수: {{서버수}}")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("호스트명", "OS종류", "제조사", "메모리용량")):
        cell.text = text
    doc.save(OUT / "서버현황_보고서.docx")


if __name__ == "__main__":
    main()
